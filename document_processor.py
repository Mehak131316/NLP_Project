import os
import re
import json
import time
import unicodedata
from collections import Counter
import numpy as np
import pandas as pd
from PIL import Image, ImageSequence

import pymupdf
import docx
from bs4 import BeautifulSoup
import nltk
from nltk.tokenize import word_tokenize, sent_tokenize
from nltk.corpus import stopwords
from nltk.stem import SnowballStemmer
from nltk.util import ngrams as nltk_ngrams
import spacy
from spacy.language import Language
from spacy.util import filter_spans

# Ensure NLTK resources
for res in ["punkt_tab", "punkt", "stopwords", "averaged_perceptron_tagger_eng", "averaged_perceptron_tagger"]:
    nltk.download(res, quiet=True)

# Load spaCy model
try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    from spacy.cli import download
    download("en_core_web_sm")
    nlp = spacy.load("en_core_web_sm")
nlp.max_length = 3_000_000

# ----------------- OCR ENGINE & RUPEE REPAIR -----------------
_OCR = {"tried": False, "engine": None, "version": "3.9.2"}

CURRENCY_BEFORE = re.compile(r"(?:₹|Rs\.?|INR|Re\.?|\$|USD|€|£)\s*$", re.IGNORECASE)
UNIT_AFTER = re.compile(r"\s*(?:shares?|units?|employees?|hours?|days?|months?|years?|sq|square|kg|km|copies|pages|words|%)\b", re.IGNORECASE)
YEN_MISREAD = re.compile(r"[¥￥](?=\s?\d)")
SLASH_DASH_NUM = re.compile(r"(?<![\d,.])\d[\d,]*(?:\.\d+)?(?=/-)")
LAKH_NUM = re.compile(r"(?<![\d,.])\d{1,2}(?:,\d{2})+,\d{3}(?:\.\d+)?(?!\d)(?!,\d)")
WESTERN_NUM = re.compile(r"(?<![\d,.])\d{1,3}(?:,\d{3})+(?:\.\d+)?(?!\d)(?!,\d)")
MONEY_CUES = re.compile(r"\b(?:stipend|salary|ctc|fees?|penalty|penalties|bonus|compensation|remuneration|amount|sum|payable|pay|paid|rent|deposit|costs?|damages|rupees)\b", re.IGNORECASE)

def get_ocr_engine():
    if _OCR["tried"]:
        return _OCR["engine"]
    _OCR["tried"] = True
    try:
        from rapidocr import RapidOCR
        _OCR["engine"] = RapidOCR(params={"Global.log_level": "warning"})
    except Exception as e:
        print(f"RapidOCR warning: {e}")
        _OCR["engine"] = None
    return _OCR["engine"]

def repair_ocr_text(text, mode="safe"):
    if mode == "off":
        return text, 0, 0
    text, n_yen = YEN_MISREAD.subn("₹", text)
    has_marker = lambda pos: bool(CURRENCY_BEFORE.search(text[max(0, pos - 6):pos]))
    inserts, unresolved = set(), 0
    for m in SLASH_DASH_NUM.finditer(text):
        if not has_marker(m.start()):
            inserts.add(m.start())
    for m in LAKH_NUM.finditer(text):
        if not has_marker(m.start()) and not UNIT_AFTER.match(text, m.end()):
            inserts.add(m.start())
    for m in WESTERN_NUM.finditer(text):
        if m.start() in inserts or has_marker(m.start()) or UNIT_AFTER.match(text, m.end()):
            continue
        if MONEY_CUES.search(text[max(0, m.start() - 40):m.start()]):
            if mode == "context":
                inserts.add(m.start())
            else:
                unresolved += 1
    for pos in sorted(inserts, reverse=True):
        text = text[:pos] + "₹" + text[pos:]
    return text, n_yen + len(inserts), unresolved

def _ocr_lines(img, engine):
    res = engine(np.asarray(img.convert("RGB")))
    if res is None or res.txts is None or len(res.txts) == 0:
        return "", 0.0
    items = []
    for box, txt, sc in zip(res.boxes, res.txts, res.scores):
        ys = [p[1] for p in box]
        xs = [p[0] for p in box]
        items.append((float(np.mean(ys)), min(xs), max(ys) - min(ys), txt, float(sc)))
    items.sort(key=lambda t: t[0])
    lines, cur, cur_y, cur_h = [], [], None, 0.0
    for y, x, h, txt, sc in items:
        if cur_y is None or abs(y - cur_y) <= 0.6 * max(cur_h, h):
            cur.append((x, txt))
            cur_y = y if cur_y is None else (cur_y + y) / 2
            cur_h = max(cur_h, h)
        else:
            lines.append(" ".join(t for _, t in sorted(cur)))
            cur, cur_y, cur_h = [(x, txt)], y, h
    if cur:
        lines.append(" ".join(t for _, t in sorted(cur)))
    return "\n".join(lines), 100 * float(np.mean([i[4] for i in items]))

def ocr_pil_image(img):
    engine = get_ocr_engine()
    if engine is None:
        return "", 0.0, 0, 0
    try:
        text, conf = _ocr_lines(img, engine)
    except Exception as e:
        return "", 0.0, 0, 0
    text, repairs, unresolved = repair_ocr_text(text)
    return text, conf, repairs, unresolved

# ----------------- REGEX & TOKEN PATTERNS -----------------
KEEP_HYPHEN = {"non", "self", "co", "ex", "anti", "work", "part", "full", "year", "long", "short",
               "third", "cross", "on", "off", "in", "out", "pre", "post", "sub", "inter", "multi"}
ABBR_END = re.compile(r"\b(?:Pvt|Ltd|Co|Inc|No|Rs|Sr|Jr|Dr|Mr|Mrs|Ms|viz|Govt|Dept)\.$")

ABBREVIATIONS = r"(?:Pvt|Ltd|Co|Inc|No|Rs|Sr|Jr|Dr|Mr|Mrs|Ms|viz|approx|Govt|Dept)\.|w\.e\.f\.|i\.e\.|e\.g\.|a\.m\.|p\.m\.|B\.Tech|M\.Tech|B\.E\.|M\.B\.A\.|etc\."
CUSTOM_PATTERNS = [
    ("MONEY",   r"(?:₹|Rs\.?|INR)\s?\d[\d,]*(?:\.\d+)?(?:/-)?(?:\s?(?:LPA|lakhs?|crores?))?"),
    ("STATUTE", r"(?:[A-Z][a-z]+\s)+Act,?\s\d{4}|(?:Section|Sec\.)\s\d+[A-Z]?(?:\(\d+\))?"),
    ("CLAUSE",  r"(?:Clause|Article)\s\d+(?:\.\d+)*(?:\([a-z]{1,3}\))?|\b\d+\.\d+(?:\.\d+)*\([a-z]{1,3}\)"),
    ("DATE",    r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b"),
    ("TIME",    r"\b\d{1,2}(?::\d{2})?\s?(?:a\.m\.|p\.m\.)"),
    ("ABBR",    ABBREVIATIONS),
    ("HYPHEN",  r"\b[A-Za-z]+(?:-[A-Za-z]+)+\b"),
    ("NUMBER",  r"\b\d+(?:,\d+)*(?:\.\d+)?\b"),
    ("WORD",    r"[A-Za-z]+(?:'[a-z]+)?"),
    ("PUNCT",   r"[^\w\s]"),
]
CUSTOM_RE = re.compile("|".join(f"(?P<{n}>{p})" for n, p in CUSTOM_PATTERNS))
PROTECTED = {"MONEY", "STATUTE", "CLAUSE", "DATE", "TIME", "ABBR", "HYPHEN"}

# Number normalization
DUR_PAREN = re.compile(r"\b(\d+)\s*\(\s*[a-z\- ]+\s*\)\s*(days?|weeks?|months?|years?)", re.IGNORECASE)
DURATION = re.compile(r"\b\d+\s+(?:days?|weeks?|months?|years?)\b", re.IGNORECASE)

# Stopwords
NLTK_STOP = set(stopwords.words("english"))
SPACY_STOP = set(nlp.Defaults.stop_words)
PROTECT = {"not", "no", "nor", "never", "without", "unless", "except", "neither", "none", "cannot",
           "n't", "shall", "may", "must", "will", "should", "can", "any", "all", "only",
           "before", "after", "during", "within", "until", "against"}
ARCHAIC = {"hereby", "herein", "hereto", "hereof", "thereof", "therein", "whereas", "witnesseth",
           "hereinafter", "aforesaid"}
CUSTOM_STOP = (NLTK_STOP - PROTECT) | ARCHAIC
FUNCTION_WORDS = NLTK_STOP | ARCHAIC
MODALS = {"shall", "may", "must", "will", "should", "can"}

stemmer_snowball = SnowballStemmer("english")

# ----------------- CUSTOM NER COMPONENT -----------------
DOMAIN_ENT_PATTERNS = [
    ("MONEY", CUSTOM_PATTERNS[0][1]),
    ("DATE", r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b"),
    ("LAW", r"(?:Section|Sec\.)\s\d+[A-Z]?(?:\(\d+\))?|(?:[A-Z][a-z]+\s)+Act,?\s\d{4}"),
    ("DURATION", r"\b\d+\s*(?:\([a-z\- ]+\)\s*)?(?:days?|weeks?|months?|years?)\b"),
    ("DESIGNATION", r"\b(?:(?:Senior|Junior|Lead|Principal|Associate|Chief)\s)?(?:(?:Software|Data|Research|Sales|Marketing|HR|Finance|Product|Technical)\s)?(?:Engineer|Developer|Analyst|Manager|Consultant|Executive|Scientist|Designer|Officer|Architect)\b"),
]

@Language.component("domain_entities_v2")
def domain_entities_v2(doc):
    new = []
    for label, pat in DOMAIN_ENT_PATTERNS:
        for m in re.finditer(pat, doc.text):
            sp = doc.char_span(m.start(), m.end(), label=label, alignment_mode="expand")
            if sp is not None:
                new.append(sp)
    kept = [e for e in doc.ents if not any(e.start < n.end and n.start < e.end for n in new)]
    doc.ents = filter_spans(new + kept)
    return doc

nlp_custom = spacy.load("en_core_web_sm")
if "domain_entities_v2" not in nlp_custom.pipe_names:
    nlp_custom.add_pipe("domain_entities_v2", after="ner")

# ----------------- POS DOMAIN LEXICON -----------------
DOMAIN_LEXICON = {
    w: "NN" for w in [
        "employee", "employer", "company", "intern", "consultant", "contractor",
        "freelancer", "worker", "candidate", "agreement", "ctc", "stipend",
        "non-compete", "non-solicitation", "gratuity", "annum"
    ]
}
DOMAIN_LEXICON.update({"w.e.f.": "IN"})
NOUN_VERB = {"notice", "bond", "sign", "release", "claim", "breach", "service"}
MONEY_TOKEN = re.compile(r"^(?:₹|Rs\.?|INR)\s?\d")

def tag_spacy_doc(words):
    doc = spacy.tokens.Doc(nlp.vocab, words=words)
    for _, proc in nlp.pipeline:
        doc = proc(doc)
    return [t.tag_ for t in doc]

def tag_rules_doc(words):
    tags = tag_spacy_doc(words)
    for i, w in enumerate(words):
        lw = w.lower()
        if MONEY_TOKEN.match(w):
            tags[i] = "CD"
        elif lw in DOMAIN_LEXICON:
            tags[i] = DOMAIN_LEXICON[lw]
        elif lw in NOUN_VERB and i > 0:
            prev = tags[i - 1]
            if prev in {"MD", "TO"} or (prev == "RB" and i > 1 and tags[i - 2] == "MD"):
                tags[i] = "VB"
            elif prev in {"DT", "JJ", "PRP$", "POS"}:
                tags[i] = "NN"
    return tags

# ----------------- CLEANING & DE-HYPHENATION -----------------
def _dehyphenate(m):
    return f"{m.group(1)}-{m.group(2)}" if m.group(1).lower() in KEEP_HYPHEN else f"{m.group(1)}{m.group(2)}"

def rejoin_wrapped_lines(lines):
    lens = sorted(len(l) for l in lines if len(l) > 20)
    width = lens[int(0.9 * (len(lens) - 1))] if lens else 0
    out, last_len = [], 0
    for l in lines:
        if out and out[-1] and l:
            prev = out[-1]
            ended = bool(re.search(r"[.:;!?]$", prev)) and not ABBR_END.search(prev)
            if not ended and (l[0].islower() or last_len >= 0.8 * width):
                out[-1] = prev + " " + l
                last_len = len(l)
                continue
        out.append(l)
        last_len = len(l)
    return out

def clean_text(text):
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[\u2018\u2019]", "'", text)
    text = re.sub(r"[\u201c\u201d]", '"', text)
    text = re.sub(r"(\w+)-\n(\w)", _dehyphenate, text)
    text = re.sub(r"(?im)^\s*page\s+\d+(\s+of\s+\d+)?\s*$", "", text)
    text = re.sub(r"(?m)^\s*\d+\s*$", "", text)
    text = re.sub(r"(?m)^\s*(?:\d+(?:\.\d+)*|[a-z]|[ivx]+)[.)]\s+", "", text)
    lines = [l.strip() for l in text.split("\n")]
    counts = Counter(l for l in lines if 0 < len(l) < 60)
    repeated = {l for l, c in counts.items() if c >= 3}
    lines = [l for l in lines if l not in repeated]
    lines = rejoin_wrapped_lines(lines)
    text = "\n".join(lines)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()

def normalise_numbers(text):
    return DUR_PAREN.sub(lambda m: f"{m.group(1)} {m.group(2)}", text)

def custom_tokenize(text, with_types=False):
    out = [(m.group(), m.lastgroup) for m in CUSTOM_RE.finditer(text)]
    return out if with_types else [t for t, _ in out]

def hybrid_doc(text):
    d = nlp.tokenizer(text)
    spans = []
    for m in CUSTOM_RE.finditer(text):
        if m.lastgroup in PROTECTED:
            sp = d.char_span(m.start(), m.end(), alignment_mode="expand")
            if sp is not None and len(sp) > 1:
                spans.append(sp)
    with d.retokenize() as retok:
        for sp in filter_spans(spans):
            retok.merge(sp)
    return d

def hybrid_tokenize(text):
    return [t.text for t in hybrid_doc(text) if not t.is_space]

def run_pipes(doc):
    for _, proc in nlp.pipeline:
        doc = proc(doc)
    return doc

def final_doc(text):
    return run_pipes(hybrid_doc(normalise_numbers(clean_text(text))))

def final_token(t):
    if " " in t.text or "-" in t.text or not t.is_alpha:
        return t.text.lower()
    return t.lemma_.lower()

def pipeline_final(text):
    doc = final_doc(text)
    return [final_token(t) for t in doc if not t.is_space and not t.is_punct and t.lemma_.lower() not in CUSTOM_STOP]

def is_meaningful(gram):
    return bool(gram and gram[0] not in FUNCTION_WORDS and gram[-1] not in FUNCTION_WORDS)

def extract_obligations(spacy_doc_obj):
    rows = []
    for tok in spacy_doc_obj:
        if tok.lower_ in MODALS and tok.head.pos_ in {"VERB", "AUX"}:
            verb = tok.head
            subj = next((c for c in verb.children if c.dep_ in {"nsubj", "nsubjpass"}), None)
            neg = any(c.dep_ == "neg" for c in verb.children)
            obj = next((c for c in verb.children if c.dep_ in {"dobj", "attr", "oprd"}), None)
            rows.append({
                "subject": " ".join(t.text for t in subj.subtree) if subj else "",
                "modal": tok.text,
                "negated": neg,
                "verb": verb.lemma_,
                "object": " ".join(t.text for t in obj.subtree) if obj else "",
                "type": ("PROHIBITION" if neg else "OBLIGATION") if tok.lower_ in {"shall", "must"}
                        else ("RIGHT" if not neg else "RESTRICTED RIGHT")
            })
    return rows

# ----------------- MAIN DOCUMENT PROCESSOR -----------------
def process_single_document(file_content, filename):
    """
    Process any single document (PDF, DOCX, TXT, HTML, Images) through the complete NLP pipeline.
    Calculates execution timing for every single stage.
    """
    t_start = time.perf_counter()
    timings = {}
    ext = os.path.splitext(filename)[1].lower()
    
    # 1. Extraction
    t0 = time.perf_counter()
    raw_text = ""
    pages_count = 1
    ocr_pages_count = 0
    ocr_confidences = []
    ocr_repairs_count = 0
    
    if ext == ".pdf":
        if isinstance(file_content, bytes):
            doc = pymupdf.open(stream=file_content, filetype="pdf")
        else:
            doc = pymupdf.open(file_content)
        pages_count = doc.page_count
        parts = []
        for pno, page in enumerate(doc):
            txt = page.get_text()
            if len(txt.strip()) < 25 and page.get_images():
                # Trigger OCR
                pix = page.get_pixmap(dpi=200, colorspace=pymupdf.csGRAY)
                img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
                ocr_txt, conf, r_rep, _ = ocr_pil_image(img)
                txt = ocr_txt
                ocr_pages_count += 1
                ocr_confidences.append(conf)
                ocr_repairs_count += r_rep
            parts.append(txt)
        raw_text = "\n".join(parts)
        doc.close()
    elif ext == ".docx":
        if isinstance(file_content, bytes):
            import io
            doc = docx.Document(io.BytesIO(file_content))
        else:
            doc = docx.Document(file_content)
        raw_text = "\n".join(p.text for p in doc.paragraphs)
    elif ext in [".html", ".htm"]:
        html_str = file_content.decode("utf-8", errors="ignore") if isinstance(file_content, bytes) else open(file_content, encoding="utf-8", errors="ignore").read()
        raw_text = BeautifulSoup(html_str, "html.parser").get_text(separator="\n")
    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp"]:
        import io
        im = Image.open(io.BytesIO(file_content)) if isinstance(file_content, bytes) else Image.open(file_content)
        ocr_txt, conf, r_rep, _ = ocr_pil_image(im)
        raw_text = ocr_txt
        ocr_pages_count = 1
        ocr_confidences.append(conf)
        ocr_repairs_count = r_rep
    else:
        raw_text = file_content.decode("utf-8", errors="ignore") if isinstance(file_content, bytes) else open(file_content, encoding="utf-8", errors="ignore").read()

    timings["extraction_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 2. Cleaning & De-hyphenation
    t0 = time.perf_counter()
    cleaned = clean_text(raw_text)
    normalized = normalise_numbers(cleaned)
    timings["cleaning_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 3. Tokenization (Built-in, Custom, Hybrid)
    t0 = time.perf_counter()
    raw_tokens = word_tokenize(cleaned)
    custom_toks = custom_tokenize(normalized)
    hybrid_toks = hybrid_tokenize(normalized)
    timings["tokenization_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 4. POS Tagging
    t0 = time.perf_counter()
    sample_words = hybrid_toks[:120] if len(hybrid_toks) > 120 else hybrid_toks
    spacy_tags = tag_spacy_doc(sample_words) if sample_words else []
    rule_tags = tag_rules_doc(sample_words) if sample_words else []
    pos_results = [{"word": w, "spacy_tag": st, "rule_tag": rt} for w, st, rt in zip(sample_words, spacy_tags, rule_tags)]
    timings["pos_tagging_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 5. Named Entity Recognition (NER)
    t0 = time.perf_counter()
    custom_spacy_doc = nlp_custom(normalized[:15000]) if normalized else nlp_custom("")
    entities = [
        {"text": e.text, "label": e.label_, "start": e.start_char, "end": e.end_char}
        for e in custom_spacy_doc.ents
    ]
    timings["ner_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 6. Lemmatization, Stopword Filtering & Pipeline Sents
    t0 = time.perf_counter()
    f_doc = final_doc(normalized[:15000]) if normalized else nlp("")
    
    final_sents = []
    for sent in f_doc.sents:
        toks = [final_token(t) for t in sent if not t.is_space and not t.is_punct]
        filtered = [t for t in toks if t not in CUSTOM_STOP]
        if filtered:
            final_sents.append(filtered)
            
    final_tokens = [t for s in final_sents for t in s]
    timings["lemmatization_stopwords_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 7. Obligations Extraction
    t0 = time.perf_counter()
    obligations_list = extract_obligations(f_doc)
    timings["obligations_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 8. N-Grams
    t0 = time.perf_counter()
    bigrams = Counter()
    trigrams = Counter()
    for s in final_sents:
        if len(s) >= 2:
            bigrams.update(nltk_ngrams(s, 2))
        if len(s) >= 3:
            trigrams.update(nltk_ngrams(s, 3))
            
    top_bigrams = [{"phrase": " ".join(g), "frequency": f, "meaningful": is_meaningful(g)} for g, f in bigrams.most_common(12)]
    top_trigrams = [{"phrase": " ".join(g), "frequency": f, "meaningful": is_meaningful(g)} for g, f in trigrams.most_common(12)]
    timings["ngrams_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    # 9. Pipeline Transformations Comparison
    t0 = time.perf_counter()
    # Pipeline A
    pipe_a_toks = [stemmer_snowball.stem(t.lower()) for t in raw_tokens if t.lower() not in NLTK_STOP and re.search(r"\w", t)]
    # Pipeline B
    pipe_b_toks = [t.lemma_.lower() for t in nlp(cleaned[:10000]) if not t.is_space and not t.is_punct]
    timings["pipelines_ms"] = round((time.perf_counter() - t0) * 1000, 2)
    
    total_ms = round((time.perf_counter() - t_start) * 1000, 2)
    timings["total_ms"] = total_ms

    return {
        "success": True,
        "filename": filename,
        "format": ext.replace(".", "").upper(),
        "pages": pages_count,
        "ocr_used": ocr_pages_count > 0,
        "ocr_pages": ocr_pages_count,
        "ocr_mean_confidence": round(float(np.mean(ocr_confidences)), 1) if ocr_confidences else None,
        "ocr_rupee_repairs": ocr_repairs_count,
        "timings": timings,
        "total_time_ms": total_ms,
        "total_time_sec": round(total_ms / 1000, 3),
        "statistics": {
            "characters": len(raw_text),
            "cleaned_characters": len(cleaned),
            "sentences_count": len(list(f_doc.sents)),
            "raw_tokens_count": len(raw_tokens),
            "hybrid_tokens_count": len(hybrid_toks),
            "final_tokens_count": len(final_tokens),
            "vocabulary_size": len(set(final_tokens))
        },
        "text_preview": cleaned[:1200] + ("..." if len(cleaned) > 1200 else ""),
        "raw_text_preview": raw_text[:800] + ("..." if len(raw_text) > 800 else ""),
        "sample_tokens": {
            "raw_nltk": raw_tokens[:40],
            "custom": custom_toks[:40],
            "hybrid": hybrid_toks[:40],
            "final": final_tokens[:40]
        },
        "pos_tags": pos_results,
        "entities": entities[:40],
        "obligations": obligations_list[:15],
        "ngrams": {
            "bigrams": top_bigrams,
            "trigrams": top_trigrams
        },
        "pipeline_comparison": {
            "Pipeline A": {
                "desc": "Cleaning -> NLTK Tokenization -> Snowball Stemming -> NLTK Stopwords",
                "tokens_count": len(pipe_a_toks),
                "sample": pipe_a_toks[:25]
            },
            "Pipeline B": {
                "desc": "Cleaning -> spaCy Tokenization -> spaCy Lemmatization -> No Stopwords Removed",
                "tokens_count": len(pipe_b_toks),
                "sample": pipe_b_toks[:25]
            },
            "Final Pipeline": {
                "desc": "Cleaning -> Number Normalization -> Hybrid Tokenization -> spaCy Lemmatization -> Domain-Aware Stopwords",
                "tokens_count": len(final_tokens),
                "sample": final_tokens[:25]
            }
        }
    }
